"""Local market analysis (#53): find businesses in an area that could become clients of the company, and
write a report into the company files.

    await start(person, args)   -> {"started": True, ...} (runs in the background, reports on Telegram)

Steps: geocode the place (Open-Meteo) -> businesses of the chosen industries within the radius from
OpenStreetMap (Overpass API) -> a quick check of each listed website (HTTPS, mobile viewport, speed,
copyright year, online booking) -> for businesses OpenStreetMap lists without a website, a limited number
of web searches (Claude + Anthropic web search) to find one -> a fit score with reasons -> Markdown summary
+ Excel list in the company share (`Market analysis/`). A repeat run for the same area marks what's new.

Rules: only public business listings; no personal data of individuals is looked up; Alfred never contacts
anyone from here (outreach is a human step, Swiss UWG/FADP). Website checks never go to private, loopback,
link-local or tailnet addresses (resolved and checked on every redirect). OpenStreetMap data © its
contributors (ODbL), credited in every report.
"""
from __future__ import annotations

import asyncio
import io
import ipaddress
import json
import logging
import math
import re
import socket
import time
from datetime import datetime
from urllib.parse import urljoin, urlsplit
from zoneinfo import ZoneInfo

import httpx

from . import db, usage
from .config import settings
from .identity.registry import Person

log = logging.getLogger(__name__)

OVERPASS_URL = "https://overpass-api.de/api/interpreter"
FOLDER = "Market analysis"
MAX_RADIUS_KM, MAX_BUSINESSES, MAX_LOOKUPS = 30, 300, 40
SITE_TIMEOUT_S, SITE_MAX_BYTES, SITE_CONCURRENCY = 8, 600_000, 6
RUN_TIMEOUT_S = 15 * 60
FIT_SCORE = 50
_SEEN_KEY = "market_analysis_seen"
_running: set[int] = set()

# Industry words (English and German) -> OpenStreetMap tags. "*" = any value of that key.
INDUSTRIES: dict[str, list[tuple[str, str]]] = {
    "restaurant": [("amenity", "restaurant")], "cafe": [("amenity", "cafe")], "bar": [("amenity", "bar"), ("amenity", "pub")],
    "fast food": [("amenity", "fast_food")], "bakery": [("shop", "bakery"), ("shop", "pastry")],
    "butcher": [("shop", "butcher")], "hairdresser": [("shop", "hairdresser")], "beauty": [("shop", "beauty")],
    "florist": [("shop", "florist")], "clothes": [("shop", "clothes")], "shoes": [("shop", "shoes")],
    "jewelry": [("shop", "jewelry")], "optician": [("shop", "optician")], "books": [("shop", "books")],
    "car repair": [("shop", "car_repair")], "car dealer": [("shop", "car")], "bicycle": [("shop", "bicycle")],
    "hardware": [("shop", "hardware"), ("shop", "doityourself")], "furniture": [("shop", "furniture")],
    "electronics": [("shop", "electronics")], "wine": [("shop", "wine"), ("craft", "winery")],
    "hotel": [("tourism", "hotel")], "guest house": [("tourism", "guest_house"), ("tourism", "apartment")],
    "dentist": [("amenity", "dentist"), ("healthcare", "dentist")], "doctor": [("amenity", "doctors")],
    "physiotherapy": [("healthcare", "physiotherapist")], "pharmacy": [("amenity", "pharmacy")],
    "veterinary": [("amenity", "veterinary")], "fitness": [("leisure", "fitness_centre")],
    "lawyer": [("office", "lawyer")], "accountant": [("office", "accountant"), ("office", "tax_advisor")],
    "architect": [("office", "architect")], "real estate": [("office", "estate_agent")],
    "insurance": [("office", "insurance")], "it": [("office", "it")], "plumber": [("craft", "plumber")],
    "electrician": [("craft", "electrician")], "carpenter": [("craft", "carpenter")], "painter": [("craft", "painter")],
    "gardener": [("craft", "gardener")], "roofer": [("craft", "roofer")],
    "photographer": [("craft", "photographer"), ("shop", "photo")], "caterer": [("craft", "caterer")],
    "driving school": [("amenity", "driving_school")], "music school": [("amenity", "music_school")],
    "shop": [("shop", "*")], "office": [("office", "*")], "craft": [("craft", "*")],
}
SYNONYMS = {
    "restaurants": "restaurant", "gastronomy": "restaurant", "gastronomie": "restaurant", "restaurants and cafes": "restaurant",
    "café": "cafe", "cafes": "cafe", "coffee": "cafe", "bars": "bar", "pub": "bar", "pubs": "bar", "takeaway": "fast food",
    "bäckerei": "bakery", "bakeries": "bakery", "metzgerei": "butcher", "coiffeur": "hairdresser", "hair salon": "hairdresser",
    "salon": "hairdresser", "hairdressers": "hairdresser", "friseur": "hairdresser", "kosmetik": "beauty", "nails": "beauty",
    "beauty salon": "beauty", "blumen": "florist", "fashion": "clothes", "mode": "clothes", "boutique": "clothes",
    "schmuck": "jewelry", "optiker": "optician", "bookshop": "books", "garage": "car repair", "autogarage": "car repair",
    "mechanic": "car repair", "car": "car dealer", "velo": "bicycle", "bike": "bicycle", "hotels": "hotel",
    "zahnarzt": "dentist", "dentists": "dentist", "arzt": "doctor", "doctors": "doctor", "physio": "physiotherapy",
    "physiotherapist": "physiotherapy", "apotheke": "pharmacy", "tierarzt": "veterinary", "vet": "veterinary",
    "gym": "fitness", "fitness studio": "fitness", "anwalt": "lawyer", "lawyers": "lawyer", "treuhand": "accountant",
    "accounting": "accountant", "architekt": "architect", "immobilien": "real estate", "versicherung": "insurance",
    "informatik": "it", "sanitär": "plumber", "plumbers": "plumber", "elektriker": "electrician", "schreiner": "carpenter",
    "maler": "painter", "gärtner": "gardener", "dachdecker": "roofer", "fotograf": "photographer", "catering": "caterer",
    "fahrschule": "driving school", "musikschule": "music school", "shops": "shop", "retail": "shop", "laden": "shop",
    "trades": "craft", "handwerk": "craft", "offices": "office",
}
BOOKING_MATTERS = {"restaurant", "cafe", "bar", "hairdresser", "beauty", "hotel", "guest house", "dentist", "doctor",
                   "physiotherapy", "veterinary", "fitness", "driving school", "music school", "car repair"}
SOCIAL = ("facebook.com", "instagram.com", "tiktok.com", "linktr.ee", "business.site", "wixsite.com/")
DIRECTORIES = ("local.ch", "search.ch", "tripadvisor", "google.", "yelp.", "booking.com", "maps.apple")


class MarketError(ValueError):
    """Message is safe to show to the person / the model."""


def resolve_industries(words: list[str]) -> tuple[list[str], list[str]]:
    known, unknown = [], []
    for w in words or []:
        key = " ".join(str(w).lower().split())
        key = SYNONYMS.get(key, key)
        if key not in INDUSTRIES and key.endswith("s") and key[:-1] in INDUSTRIES:
            key = key[:-1]
        (known if key in INDUSTRIES else unknown).append(key)
    return sorted(set(known)), unknown


# --- data ---------------------------------------------------------------------------------------

def _overpass_query(tags: list[tuple[str, str]], lat: float, lon: float, radius_m: int) -> str:
    def sel(k: str, v: str) -> str:
        return f'["{k}"]' if v == "*" else f'["{k}"="{v}"]'
    parts = "".join(f'nwr{sel(k, v)}["name"](around:{radius_m},{lat:.5f},{lon:.5f});' for k, v in tags)
    return f"[out:json][timeout:90];({parts});out center tags {MAX_BUSINESSES * 3};"


def _distance_km(lat1, lon1, lat2, lon2) -> float:
    p = math.pi / 180
    a = (math.sin((lat2 - lat1) * p / 2) ** 2
         + math.cos(lat1 * p) * math.cos(lat2 * p) * math.sin((lon2 - lon1) * p / 2) ** 2)
    return 12742 * math.asin(math.sqrt(a))


def _norm_url(url: str | None) -> str | None:
    url = (url or "").strip().split(";")[0].strip()
    if not url:
        return None
    if not re.match(r"(?i)^https?://", url):
        url = "https://" + url.lstrip("/")
    return url


def _category(tags: dict, industries: list[str]) -> str:
    for ind in industries:
        for k, v in INDUSTRIES[ind]:
            if k in tags and (v == "*" or tags[k] == v):
                return ind if v != "*" else f"{ind}: {tags[k].replace('_', ' ')}"
    return "other"


def parse_businesses(data: dict, industries: list[str], lat: float, lon: float) -> list[dict]:
    out, seen = [], set()
    for el in data.get("elements", []):
        tags = el.get("tags") or {}
        name = (tags.get("name") or "").strip()
        la, lo = el.get("lat") or (el.get("center") or {}).get("lat"), el.get("lon") or (el.get("center") or {}).get("lon")
        if not name or la is None:
            continue
        street = " ".join(x for x in (tags.get("addr:street"), tags.get("addr:housenumber")) if x)
        town = tags.get("addr:city") or tags.get("addr:place") or ""
        key = (name.lower(), street.lower())
        if key in seen:
            continue
        seen.add(key)
        out.append({
            "osm_id": f"{el.get('type')}/{el.get('id')}", "name": name, "category": _category(tags, industries),
            "street": street, "postcode": tags.get("addr:postcode", ""), "town": town,
            "distance_km": round(_distance_km(lat, lon, la, lo), 2),
            "website": _norm_url(tags.get("website") or tags.get("contact:website") or tags.get("url")),
            "phone": tags.get("phone") or tags.get("contact:phone") or "",
            "email": tags.get("email") or tags.get("contact:email") or "",
            "opening_hours": tags.get("opening_hours", ""), "website_source": "OpenStreetMap",
        })
    out.sort(key=lambda b: b["distance_km"])
    return out[:MAX_BUSINESSES]


async def fetch_businesses(industries: list[str], lat: float, lon: float, radius_km: float) -> list[dict]:
    tags = [t for ind in industries for t in INDUSTRIES[ind]]
    query = _overpass_query(tags, lat, lon, int(radius_km * 1000))
    headers = {"User-Agent": f"{settings.assistant_name}-assistant/1.0 (market analysis; {settings.contact_email})"}
    try:
        async with httpx.AsyncClient(timeout=120) as c:
            r = await c.post(OVERPASS_URL, data={"data": query}, headers=headers)
            r.raise_for_status()
            data = r.json()
    except (httpx.HTTPError, ValueError) as e:
        raise MarketError(f"OpenStreetMap couldn't be queried right now ({type(e).__name__}). Try again later.")
    usage.record_request("openstreetmap", "overpass")
    return parse_businesses(data, industries, lat, lon)


# --- website checks -----------------------------------------------------------------------------

async def _public_host(host: str) -> bool:
    """Every address the host resolves to must be public (no LAN, loopback, link-local, tailnet/CGNAT)."""
    if not host:
        return False
    try:
        infos = await asyncio.get_running_loop().getaddrinfo(host, None, type=socket.SOCK_STREAM)
    except OSError:
        return False
    ips = {ipaddress.ip_address(i[4][0]) for i in infos}
    return bool(ips) and all(ip.is_global for ip in ips)


YEAR = re.compile(r"(?:©|&copy;|copyright)\D{0,25}((?:19|20)\d{2})(?:\s*(?:-|–|&ndash;|bis)\s*((?:19|20)\d{2}))?", re.I)
BOOKING = re.compile(r"\b(book(?:ing)?|reserv|termin|buchen|online[- ]?shop|order online|bestellen|calendly|"
                     r"timify|treatwell|opentable|thefork|quandoo)\w*", re.I)


async def check_site(url: str, client: httpx.AsyncClient | None = None) -> dict:
    out = {"reachable": False, "https": False, "mobile": None, "seconds": None, "copyright_year": None,
           "booking": False, "final_url": url, "error": ""}
    own = client is None
    client = client or httpx.AsyncClient(timeout=SITE_TIMEOUT_S, follow_redirects=False,
                                         headers={"User-Agent": "Mozilla/5.0 (compatible; site check)"})
    started = time.monotonic()
    try:
        current = url
        for _ in range(5):                                   # redirects, each one checked
            parts = urlsplit(current)
            if parts.scheme not in ("http", "https") or not await _public_host(parts.hostname or ""):
                out["error"] = "not a public web address"
                return out
            async with client.stream("GET", current) as r:
                if r.is_redirect and r.headers.get("location"):
                    current = urljoin(current, r.headers["location"])
                    continue
                body = b""
                async for chunk in r.aiter_bytes():
                    body += chunk
                    if len(body) >= SITE_MAX_BYTES:
                        break
                out.update(final_url=current, reachable=r.status_code < 400, https=current.startswith("https://"))
                if r.status_code >= 400:
                    out["error"] = f"HTTP {r.status_code}"
                    return out
                html = body.decode("utf-8", errors="replace")
                out["mobile"] = bool(re.search(r'<meta[^>]+name=["\']?viewport', html, re.I))
                years = [int(y) for m in YEAR.finditer(html) for y in m.groups() if y]
                out["copyright_year"] = max(years) if years else None
                out["booking"] = bool(BOOKING.search(html))
                return out
        out["error"] = "too many redirects"
    except httpx.HTTPError as e:
        out["error"] = type(e).__name__
    finally:
        out["seconds"] = round(time.monotonic() - started, 1)
        if own:
            await client.aclose()
    return out


async def check_sites(businesses: list[dict]) -> None:
    sem = asyncio.Semaphore(SITE_CONCURRENCY)
    async with httpx.AsyncClient(timeout=SITE_TIMEOUT_S, follow_redirects=False,
                                 headers={"User-Agent": "Mozilla/5.0 (compatible; site check)"}) as client:
        async def one(b):
            if not b["website"] or _is_social(b["website"]):
                return
            async with sem:
                b["site"] = await check_site(b["website"], client)
        await asyncio.gather(*(one(b) for b in businesses))


def _is_social(url: str) -> bool:
    return any(s in url.lower() for s in SOCIAL)


# --- web search for missing websites ------------------------------------------------------------

async def find_missing_websites(businesses: list[dict], limit: int, person: Person) -> tuple[int, int]:
    """Ask Claude with web search for the official website of businesses OpenStreetMap lists without one."""
    if not settings.web_search_enabled or limit <= 0:
        return 0, 0
    from . import ai_config
    if not ai_config.configured():
        return 0, 0
    from .agent_core import client
    todo = [b for b in businesses if not b["website"]][:limit]
    for b in todo:
        b["looked_up"] = True
    found = 0
    for i in range(0, len(todo), 8):
        batch = todo[i:i + 8]
        listing = [{"id": n, "name": b["name"], "category": b["category"], "address": f"{b['street']}, {b['town']}".strip(", ")}
                   for n, b in enumerate(batch)]
        system = ("For each business, search the web for its own official website. Directory or map listings "
                  "(local.ch, search.ch, Google Maps, TripAdvisor) don't count; a Facebook or Instagram page counts "
                  "only if it is all they have. Reply with JSON only: "
                  '[{"id": 0, "website": "https://..." or null}]')
        try:
            resp = await client().messages.create(
                model=settings.claude_model, max_tokens=1500, system=system,
                tools=[{"type": "web_search_20250305", "name": "web_search", "max_uses": len(batch)}],
                messages=[{"role": "user", "content": json.dumps(listing, ensure_ascii=False)}])
        except Exception as e:  # noqa: BLE001 — a failed lookup only means fewer websites found
            log.warning("Market analysis web lookup failed: %s", e)
            continue
        usage.record_claude(getattr(resp, "usage", None), kind="market_analysis", user_id=person.id)
        text = "".join(getattr(b, "text", "") for b in resp.content if getattr(b, "type", "") == "text")
        m = re.search(r"\[.*\]", text, re.S)
        try:
            rows = json.loads(m.group(0)) if m else []
        except ValueError:
            rows = []
        for row in rows:
            try:
                b = batch[int(row.get("id"))]
            except (TypeError, ValueError, IndexError):
                continue
            url = _norm_url(row.get("website"))
            if url and not any(d in url.lower() for d in DIRECTORIES):
                b["website"], b["website_source"] = url, "web search"
                found += 1
    return len(todo), found


# --- scoring and report -------------------------------------------------------------------------

def score(b: dict, year_now: int) -> tuple[int, list[str]]:
    if not b["website"]:
        if b.get("looked_up"):
            return 90, ["no website found (web search)"]
        return 70, ["no website listed on OpenStreetMap (not looked up)"]
    if _is_social(b["website"]):
        return 80, ["only a social media page"]
    site = b.get("site") or {}
    if not site.get("reachable"):
        return 85, [f"website not reachable ({site.get('error') or 'no answer'})"]
    s, why = 20, []
    if not site.get("https"):
        s += 20; why.append("no HTTPS")
    if site.get("mobile") is False:
        s += 30; why.append("not mobile-friendly")
    if (site.get("seconds") or 0) > 3:
        s += 10; why.append(f"slow ({site['seconds']} s)")
    y = site.get("copyright_year")
    if y and y <= year_now - 3:
        s += 20; why.append(f"looks outdated (© {y})")
    if b["category"].split(":")[0] in BOOKING_MATTERS and not site.get("booking"):
        s += 10; why.append("no online booking or ordering")
    return min(s, 95), why or ["website looks fine"]


def _web_status(b: dict) -> str:
    if not b["website"]:
        return "none"
    if _is_social(b["website"]):
        return "social media only"
    site = b.get("site") or {}
    return "reachable" if site.get("reachable") else "not reachable"


def build_markdown(run: dict, businesses: list[dict]) -> str:
    fit = [b for b in businesses if b["fit"]]
    towns: dict[str, int] = {}
    for b in fit:
        towns[b["town"] or "unknown"] = towns.get(b["town"] or "unknown", 0) + 1
    lines = [f"# Market analysis: {', '.join(run['industries'])} within {run['radius_km']:g} km of {run['area']}", "",
             f"*{run['date']} · prepared by {settings.assistant_name} for {settings.company_name}*", "",
             "## Summary", "",
             f"- Businesses found: **{len(businesses)}** (OpenStreetMap)",
             f"- Fitting potential clients (score ≥ {FIT_SCORE}): **{len(fit)}**",
             f"- No website, confirmed by web search: **{sum(1 for b in businesses if not b['website'] and b.get('looked_up'))}**",
             f"- No website listed on OpenStreetMap, not looked up yet: {sum(1 for b in businesses if not b['website'] and not b.get('looked_up'))}"
             " (ask for more web lookups to check them)",
             f"- Websites found by web search: {run['found_by_search']} (of {run['lookups']} lookups)"]
    if run.get("previous"):
        lines.append(f"- New since the last report for this area: **{sum(1 for b in businesses if b.get('new'))}**")
    lines += ["", "## Shortlist", "", "| # | Business | Category | Town | Website | Score | Why |", "|---|---|---|---|---|---|---|"]
    for n, b in enumerate(fit[:40], 1):
        new = " 🆕" if b.get("new") else ""
        lines.append(f"| {n} | {b['name']}{new} | {b['category']} | {b['town'] or '–'} | {_web_status(b)} | {b['score']} | "
                     f"{'; '.join(b['reasons'])} |")
    if len(fit) > 40:
        lines.append(f"\n…and {len(fit) - 40} more in the Excel file.")
    lines += ["", "## Where they are", "", "| Town | Fitting businesses |", "|---|---|"]
    lines += [f"| {t} | {n} |" for t, n in sorted(towns.items(), key=lambda x: -x[1])]
    lines += ["", "## Method", "",
              f"Businesses tagged as {', '.join(run['industries'])} within {run['radius_km']:g} km of {run['area']} "
              "were taken from OpenStreetMap. Listed websites were opened once to check HTTPS, a mobile layout, "
              "loading time, the copyright year and online booking; for businesses without a listed website a "
              "limited number of web searches looked for one. Scores: no website (confirmed) 90, no website listed "
              "but not looked up 70, social media only 80, "
              "unreachable 85, otherwise 20 plus points per weakness (not mobile-friendly 30, outdated 20, "
              "no HTTPS 20, slow 10, no online booking 10).", "",
              "Only public business listings were used. Contacting businesses is up to you: in Switzerland, "
              "unsolicited advertising must respect the UWG (e.g. the 'no advertising' asterisk in the phone "
              "directory) and the FADP.", "",
              "Data © OpenStreetMap contributors (ODbL), https://www.openstreetmap.org/copyright", ""]
    return "\n".join(lines)


def build_excel(run: dict, businesses: list[dict]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Businesses"
    cols = ["Fit", "Score", "Why", "Name", "Category", "Street", "Postcode", "Town", "Distance km", "Website",
            "Website from", "Website status", "HTTPS", "Mobile", "Load s", "© year", "Online booking", "Phone",
            "Email", "Opening hours", "New", "OpenStreetMap"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for b in businesses:
        s = b.get("site") or {}
        ws.append(["yes" if b["fit"] else "", b["score"], "; ".join(b["reasons"]), b["name"], b["category"], b["street"],
                   b["postcode"], b["town"], b["distance_km"], b["website"] or "", b["website_source"] if b["website"] else "",
                   _web_status(b), {True: "yes", False: "no"}.get(s.get("https"), ""),
                   {True: "yes", False: "no"}.get(s.get("mobile"), ""), s.get("seconds") or "", s.get("copyright_year") or "",
                   {True: "yes", False: "no"}.get(s.get("booking"), "") if s else "", b["phone"], b["email"],
                   b["opening_hours"], "yes" if b.get("new") else "", f"https://www.openstreetmap.org/{b['osm_id']}"])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFGHIJKLMNOPQRSTUV", (5, 7, 45, 32, 18, 26, 9, 16, 11, 36, 13, 16, 7, 8, 7, 7, 9, 18, 26, 24, 5, 40)):
        ws.column_dimensions[col].width = width
    info = wb.create_sheet("About")
    for row in (("Area", run["area"]), ("Radius km", run["radius_km"]), ("Industries", ", ".join(run["industries"])),
                ("Date", run["date"]), ("Businesses", len(businesses)), ("Fitting", sum(1 for b in businesses if b["fit"])),
                ("Source", "© OpenStreetMap contributors (ODbL); web search for missing websites")):
        info.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- run ----------------------------------------------------------------------------------------

def _run_key(area: str, radius_km: float, industries: list[str]) -> str:
    return f"{' '.join(area.lower().split())}|{radius_km:g}|{','.join(industries)}"


async def analyse(person: Person, area: str, radius_km: float, industries: list[str], lookups: int,
                  check_websites: bool = True) -> dict:
    """The whole analysis; returns the saved file paths and counts (raises MarketError)."""
    from . import briefing, nas_files
    try:
        place = await briefing.geocode(area)
    except briefing.BriefingError as e:
        raise MarketError(str(e).replace("weather", "analysis"))
    businesses = await fetch_businesses(industries, place["lat"], place["lon"], radius_km)
    if not businesses:
        raise MarketError(f"OpenStreetMap lists no {', '.join(industries)} within {radius_km:g} km of {place['location']}.")
    searched, found = await find_missing_websites(businesses, lookups, person)
    if check_websites:
        await check_sites(businesses)
    year = datetime.now(ZoneInfo(settings.timezone)).year
    for b in businesses:
        b["score"], b["reasons"] = score(b, year)
        b["fit"] = b["score"] >= FIT_SCORE
    seen = db.get_setting(_SEEN_KEY, {}) or {}
    key = _run_key(place["location"], radius_km, industries)
    previous = set(seen.get(key, []))
    for b in businesses:
        b["new"] = bool(previous) and b["osm_id"] not in previous
    seen[key] = [b["osm_id"] for b in businesses]
    db.set_setting(_SEEN_KEY, dict(list(seen.items())[-50:]))
    businesses.sort(key=lambda b: (-b["score"], b["distance_km"]))
    date = datetime.now(ZoneInfo(settings.timezone)).strftime("%Y-%m-%d")
    run = {"area": place["location"], "radius_km": radius_km, "industries": industries, "date": date,
           "found_by_search": found, "lookups": searched, "previous": bool(previous)}
    stem = f"{date} {place['location'].split(',')[0]} {radius_km:g}km {' '.join(industries)}"[:120]
    md = await asyncio.to_thread(nas_files.upload, person, "company", FOLDER, f"{stem}.md",
                                 build_markdown(run, businesses).encode("utf-8"))
    xl = await asyncio.to_thread(nas_files.upload, person, "company", FOLDER, f"{stem}.xlsx", build_excel(run, businesses))
    fit = [b for b in businesses if b["fit"]]
    db.log_activity("system", f"Market analysis by {person.display_name}: {len(businesses)} businesses, {len(fit)} fitting "
                              f"({', '.join(industries)}, {radius_km:g} km of {place['location']})", user_id=person.id)
    return {"area": place["location"], "businesses": len(businesses), "fitting": len(fit),
            "no_website": sum(1 for b in businesses if not b["website"]), "new": sum(1 for b in businesses if b.get("new")),
            "top": [f"{b['name']} ({b['town'] or b['category']}): {'; '.join(b['reasons'])}" for b in fit[:5]],
            "files": [md["saved"], xl["saved"]]}


async def start(person: Person, area: str, radius_km: float = 5, industries: list[str] | None = None,
                web_lookups: int = 15, check_websites: bool = True) -> dict:
    from . import nas
    nas.check_scope(person, "company")
    if not nas.enabled():
        raise MarketError("The NAS isn't connected, so there's nowhere to save the report.")
    area = " ".join(str(area or "").split())[:80]
    if not area:
        raise MarketError("Which place should the analysis be around?")
    try:
        radius_km = float(radius_km)
    except (TypeError, ValueError):
        raise MarketError("The radius must be a number of kilometres.")
    if not 0.5 <= radius_km <= MAX_RADIUS_KM:
        raise MarketError(f"Pick a radius between 0.5 and {MAX_RADIUS_KM} km.")
    known, unknown = resolve_industries(industries or [])
    if unknown:
        raise MarketError(f"I can't search for: {', '.join(unknown)}. Known industries: {', '.join(sorted(INDUSTRIES))}.")
    if not known:
        raise MarketError("Which industries should I look at (e.g. restaurants, hairdressers, garages)?")
    lookups = max(0, min(int(web_lookups or 0), MAX_LOOKUPS))
    if person.id in _running:
        raise MarketError("An analysis for you is already running; I'll report when it's done.")
    _running.add(person.id)

    async def job():
        try:
            res = await asyncio.wait_for(analyse(person, area, radius_km, known, lookups, check_websites), RUN_TIMEOUT_S)
            text = (f"📊 Market analysis ready: {res['fitting']} of {res['businesses']} businesses look like fitting "
                    f"clients ({', '.join(known)} within {radius_km:g} km of {res['area']}; {res['no_website']} without "
                    f"a website" + (f", {res['new']} new since last time" if res["new"] else "") + ").\n\n"
                    + "\n".join(f"• {t}" for t in res["top"]) + "\n\nSaved in the company files:\n"
                    + "\n".join(f"• {f}" for f in res["files"]))
        except MarketError as e:
            text = f"The market analysis couldn't finish: {e}"
        except asyncio.TimeoutError:
            text = "The market analysis took too long and was stopped. Try a smaller radius or fewer industries."
        except Exception:  # noqa: BLE001 — reported, never crashes the server
            log.exception("Market analysis failed")
            text = "The market analysis failed unexpectedly; the details are in the server log."
        finally:
            _running.discard(person.id)
        try:
            from .telegram_handler import send_to_user
            await send_to_user(person.id, text, note="market analysis", silent=False)
        except Exception:  # noqa: BLE001
            log.exception("Couldn't send the market analysis result")

    asyncio.get_running_loop().create_task(job())
    return {"started": True, "industries": known, "radius_km": radius_km, "area": area, "web_lookups": lookups,
            "note": f"Running in the background (usually 1-5 minutes). The report goes to the company files under "
                    f"'{FOLDER}', and the person gets a Telegram message when it's ready. Tell them that briefly."}
