"""A local stand-in for the NAS (SMB) and a copy of the Company to-do layout: used by the tests and by
tests/devserver.py, so neither ever touches the real NAS."""
from pathlib import Path


class FakeSMB:
    """smbclient stand-in: \\nas\Share\path -> <tmp>/Share/path. Shares that are 'denied' raise like TrueNAS."""

    def __init__(self, root, denied=()):
        self.root, self.denied = Path(root), set(denied)

    def _p(self, unc):
        parts = [x for x in unc.split("\\") if x]
        if parts[1] in self.denied:
            raise OSError("[NtStatus 0xc0000022] STATUS_ACCESS_DENIED")
        return self.root.joinpath(*parts[1:])

    def register_session(self, *a, **kw): pass
    def reset_connection_cache(self): pass

    def open_file(self, unc, mode="r", encoding=None, newline=None, **kw):
        path = self._p(unc)
        if "w" in mode:
            path.parent.mkdir(parents=True, exist_ok=True)
        return open(path, mode, encoding=encoding if "b" not in mode else None,
                    newline=newline if "b" not in mode else None)

    def stat(self, unc):
        path = self._p(unc)
        if not path.exists():
            raise FileNotFoundError(unc)
        return path.stat()

    def makedirs(self, unc, exist_ok=False):
        self._p(unc).mkdir(parents=True, exist_ok=exist_ok)

    def listdir(self, unc):
        return [x.name for x in self._p(unc).iterdir()]

    def scandir(self, unc):
        import os
        return list(os.scandir(self._p(unc)))

    def walk(self, unc):
        import os
        base = self._p(unc)
        for dirpath, dirs, files in os.walk(base):
            rel = Path(dirpath).relative_to(base)
            yield (unc + ("\\" + str(rel).replace("/", "\\") if str(rel) != "." else "")), dirs, files

    def rename(self, a, b):
        self._p(a).rename(self._p(b))

    def remove(self, unc):
        self._p(unc).unlink()


def todo_workbook(path):
    """A copy of the Company to-do layout (3 sections, rows 6-30)."""
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabelle1"
    ws["A1"] = "COMPANY To Do Liste"
    ws["A4"], ws["H4"], ws["O4"] = "AKTUELLES", "ZUKUNFT", "FERTIG"
    for col, head in zip("ABCDHIJKOPQ", ["Position", "wichtigkeit", "Status", "beschreibung"] * 2 + ["Position", "wichtigkeit", "beschreibung"]):
        ws[f"{col}5"] = head
    for r in range(6, 31):
        for pos, imp, st in (("A", "B", "C"), ("H", "I", "J")):
            ws[f"{pos}{r}"] = 1 if r == 6 else f"=SUM({pos}{r - 1} + 1)"
            ws[f"{imp}{r}"] = 1
            ws[f"{st}{r}"] = "not started"
        ws[f"O{r}"] = 1 if r == 6 else f"=SUM(O{r - 1} + 1)"
        ws[f"P{r}"] = 1
    ws["B6"], ws["C6"], ws["D6"] = 5, "in progress", "KLG Gründung Abschliessen"
    ws["B7"], ws["C7"], ws["D7"] = 4, "in progress", "Offizielles eröffnungsdatum definieren"
    ws["B8"], ws["D8"] = 4, "Angebote erweitern"
    ws["I6"], ws["K6"] = 4, "Planung des Online Shops Starten"
    ws.conditional_formatting.add("C6:C30", CellIsRule(operator="equal", formula=['"done"'],
                                                       fill=PatternFill("solid", start_color="00FF00")))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
